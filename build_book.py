"""Writes the workbook: one Sheet1 in the shape of the hand-made book.

One block per expiry, exactly like the trading sheet: a STRIKE / QTY / PREMIUM ladder with a
'CP IF <spot>' column every 100 points, the CE side, then the PE side, then the futures row and
the block TOTAL. QTY and PREMIUM are plain numbers written from the screenshot and the NSE chain
(so the sheet opens with the book already priced), and every CP IF cell is the same IF formula
the hand-made sheet uses, so editing a quantity or a rate still recalculates the whole grid.

A manual page break sits under each block, so printing gives the current week on one page and
the next week on the next.
"""
from __future__ import annotations

import calendar
import datetime as dt
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=11, bold=True)
CENTER = Alignment(horizontal="center")

WIDTH_STRIKE = 9.57
WIDTH_QTY = 15.57
WIDTH_PREMIUM = 10.71
WIDTH_SCENARIO = 12.0

CE_FORMULA = '=IF(($A{r}>{c}${a}),($C{r}*$B{r}),($A{r}-{c}${a})*$B{r}+($B{r}*$C{r}))'
PE_FORMULA = '=IF(({c}${a}<$A{r}),(-$A{r}+{c}${a})*$B{r}+($C{r}*$B{r}),($C{r}*$B{r}))'
FUT_FORMULA = '=+({c}{a}-$C{r})*$B{r}'

MONTH_ABBR = {calendar.month_abbr[m].lower(): m for m in range(1, 13)}


def ladder_strikes(spot: float, step: int, span: float) -> list[float]:
    """Strikes from spot-span to spot+span on the chosen 50/100 grid."""
    lo = math.floor((spot - span) / step) * step
    hi = math.ceil((spot + span) / step) * step
    out, k = [], lo
    while k <= hi + 1e-9:
        out.append(float(k))
        k += step
    return out


def ordinal(day: int) -> str:
    if 11 <= day % 100 <= 13:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
    return f"{day}{suffix}"


def expiry_label(expiry: str) -> str:
    """'19-Oct-2026' -> '19th October', the way the hand-made sheet titles a block."""
    parts = (expiry or "").split("-")
    if len(parts) == 3 and parts[1].lower() in MONTH_ABBR:
        return f"{ordinal(int(parts[0]))} {calendar.month_name[MONTH_ABBR[parts[1].lower()]]}"
    return expiry or ""


def parse_expiry(expiry: str) -> dt.date | None:
    try:
        day, mon, year = expiry.split("-")
        return dt.date(int(year), MONTH_ABBR[mon.lower()], int(day))
    except (AttributeError, ValueError, KeyError):
        return None


def leg_rates(chain: dict) -> dict[float, dict[str, float]]:
    """strike -> {CE: last, PE: last} for one expiry's quoted chain."""
    out: dict[float, dict[str, float]] = {}
    for raw, legs in (chain or {}).get("strikes", {}).items():
        try:
            key = round(float(raw), 2)
        except (TypeError, ValueError):
            continue
        out[key] = {side: (leg or {}).get("last") or 0 for side, leg in (legs or {}).items()}
    return out


def _split_book(positions: list[dict]) -> tuple[dict[tuple[str, float, str], int], int, float | None,
                                                dict[tuple[str, float, str], float]]:
    """(expiry, strike, type) -> signed units, the futures quantity/price, and the terminal's own rates.

    QTY in this grid is signed the way the IF formulas need it: a Sell is positive (that is the
    side that collects premium), a Buy negative - the terminal shows the same size with the sign
    reversed. The futures row is priced as (scenario spot - entry) x qty, so it keeps the
    terminal's own sign instead.
    """
    legs: dict[tuple[str, float, str], int] = {}
    cost: dict[tuple[str, float, str], float] = {}
    futures = 0
    fut_price: float | None = None
    for p in positions:
        qty = int(p.get("qty") or 0)
        if not qty:
            continue
        sign = 1 if (p.get("side") or "Sell").lower().startswith("s") else -1
        kind = (p.get("otype") or "").upper()
        strike = round(float(p.get("strike") or 0), 2)
        if kind in ("CE", "PE") and strike:
            key = (p.get("expiry") or "", strike, kind)
            legs[key] = legs.get(key, 0) + sign * qty
            traded = [(p.get("sell_qty") or 0, p.get("sell_avg")), (p.get("buy_qty") or 0, p.get("buy_avg"))]
            traded = [(q, r) for q, r in traded if q and r]
            if traded:
                cost[key] = sum(q * float(r) for q, r in traded) / sum(q for q, _ in traded)
        elif not strike:
            # only a row with no strike is a future; an option whose type OCR lost is handled in
            # parse_grid, and must never be dumped here where it would pollute the futures line
            futures -= sign * qty
            fut_price = fut_price or p.get("mkt_price") or p.get("sell_avg") or p.get("buy_avg")
    return legs, futures, fut_price, cost


def _style_row(ws, row: int, last_col: int) -> None:
    for col in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = BOLD
        cell.border = BOX


def _round(value) -> int:
    """The sheet is hand-typed in round figures, so every rate goes in as a whole number."""
    try:
        return int(round(float(value or 0)))
    except (TypeError, ValueError):
        return 0


def _section(ws, row: int, scenarios: list[float], kind: str, ladder: list[float],
             legs: dict[tuple[str, float, str], int], rates: dict[float, dict[str, float]],
             cost: dict[tuple[str, float, str], float], expiry: str) -> tuple[int, int, int]:
    """Write one STRIKE / QTY / PREMIUM section over the full ladder and return (header, anchor, total).

    Every strike on the ladder gets a row - including the ones with no position, which come out as
    QTY 0 / PREMIUM 0. That is deliberate: it matches the hand-made sheet and leaves a blank slot to
    type into when a leg is added or changed by hand.
    """
    anchor = row + 1
    top = anchor + 1
    for col, title in enumerate(("STRIKE", "QTY", "PREMIUM"), start=1):
        ws.cell(row=row, column=col, value=title)
    for j, level in enumerate(scenarios):
        ws.cell(row=row, column=4 + j, value=f"CP IF {level:g}")
        ws.cell(row=anchor, column=4 + j, value=level)
    _style_row(ws, row, 3 + len(scenarios))
    _style_row(ws, anchor, 3 + len(scenarios))

    for i, strike in enumerate(ladder):
        r = top + i
        qty = legs.get((expiry, strike, kind), 0)
        # NSE can have no quote for a strike the book holds, so fall back to the terminal's own
        # rate for it rather than pricing a real position at zero.
        premium = _round((rates.get(strike, {}).get(kind)
                          or cost.get((expiry, strike, kind), 0)) if qty else 0)
        ws.cell(row=r, column=1, value=strike)
        ws.cell(row=r, column=2, value=qty)
        ws.cell(row=r, column=3, value=premium)
        for j in range(len(scenarios)):
            col = get_column_letter(4 + j)
            ws.cell(row=r, column=4 + j, value=(
                CE_FORMULA if kind == "CE" else PE_FORMULA).format(r=r, a=anchor, c=col))
        _style_row(ws, r, 3 + len(scenarios))

    last = top + len(ladder) - 1
    total = last + (2 if kind == "CE" else 1)      # the CE block leaves a blank row above its total
    ws.cell(row=total, column=2, value=f"=SUM(B{top}:B{last})")
    for j in range(len(scenarios)):
        col = get_column_letter(4 + j)
        ws.cell(row=total, column=4 + j, value=f"=SUM({col}{top}:{col}{last})")
    _style_row(ws, total, 3 + len(scenarios))
    return row, anchor, total


def _block(ws, first: int, positions: list[dict], chain: dict, cfg: dict, expiry: str) -> dict | None:
    """One expiry = one printed page: heading, full CE ladder, full PE ladder, futures, TOTAL.

    The ladder covers every strike from spot-span to spot+span on the 50/100 grid, plus any strike
    the book actually holds outside that window. Strikes with no position still get a row (QTY 0,
    PREMIUM 0), exactly like the hand-made sheet, so the grid is a template to edit by hand.
    """
    spot = (chain or {}).get("spot") or cfg.get("spot") or 0
    step, span = cfg["step"], cfg["span"]
    legs, fut_qty, fut_price, cost = _split_book(positions)
    rates = leg_rates(chain)
    scenarios = ladder_strikes(spot, step * 2, span)
    held = {kind: {strike for (exp, strike, k) in legs
                   if exp == expiry and k == kind and legs[(exp, strike, k)]} for kind in ("CE", "PE")}
    if not (held["CE"] or held["PE"] or fut_qty):
        return None
    # full ladder for both sections = the spot-centred grid, extended to any held strike outside it
    ladder = sorted(set(ladder_strikes(spot, step, span)) | held["CE"] | held["PE"])

    ws.cell(row=first, column=1, value=cfg["symbol"])
    ws.cell(row=first, column=3, value=_round(spot))
    ws.cell(row=first, column=10, value=expiry_label(expiry)).alignment = CENTER
    ws.cell(row=first + 1, column=1, value="DATE:- ")
    date_cell = ws.cell(row=first + 1, column=2, value=dt.date.today())
    date_cell.number_format = "d-mmm-yy"
    ws.cell(row=first + 2, column=1, value=f"NET BOOK - {expiry}")
    for r in (first, first + 1, first + 2):
        for c in (1, 2, 3, 10):
            ws.cell(row=r, column=c).font = TITLE

    totals: list[int] = []
    anchor_row = None
    row = first + 3
    for kind in ("CE", "PE"):
        _hdr, anchor, total = _section(ws, row, scenarios, kind, ladder, legs, rates, cost, expiry)
        anchor_row = anchor
        totals.append(total)
        row = total + 2

    if fut_qty:
        ws.cell(row=row, column=1, value=cfg["symbol"])
        ws.cell(row=row, column=2, value=fut_qty)
        ws.cell(row=row, column=3, value=_round(fut_price or 0))
        for j, level in enumerate(scenarios):
            col = get_column_letter(4 + j)
            ws.cell(row=row, column=4 + j, value=(
                FUT_FORMULA.format(a=anchor_row, r=row, c=col) if anchor_row
                else f"=+({level}-$C{row})*$B{row}"))
        totals.append(row)
        _style_row(ws, row, 3 + len(scenarios))
        row += 2

    grand = row
    ws.cell(row=grand, column=3, value="TOTAL")
    for j in range(len(scenarios)):
        col = get_column_letter(4 + j)
        ws.cell(row=grand, column=4 + j, value="=" + "+".join(f"{col}{t}" for t in totals))
    _style_row(ws, grand, 3 + len(scenarios))

    rows = 2 * len(ladder) + (1 if fut_qty else 0)
    return {"first": first, "last": grand, "rows": rows, "scenarios": len(scenarios), "spot": spot}


def resolve_option_types(positions: list[dict], chains: dict[str, dict]) -> list[dict]:
    """Recover an Option Type the OCR lost, using the live NSE chain.

    parse_grid infers from a sibling leg at the same strike; when there is no sibling (a lone
    position, or a highlighted row whose type cell OCR skipped) the type stays blank and the row
    would silently vanish from the book. A row with a strike is always an option, so pick the leg
    whose last price is closest to the terminal's own market price for it.
    """
    rates_by_exp = {e: leg_rates(c) for e, c in (chains or {}).items()}
    for p in positions:
        if not p.get("strike") or (p.get("otype") or "").upper() in ("CE", "PE"):
            continue
        leg = rates_by_exp.get(p.get("expiry") or "", {}).get(round(float(p["strike"]), 2)) or {}
        ref = p.get("mkt_price") or p.get("sell_avg") or p.get("buy_avg")
        ce, pe = leg.get("CE") or 0, leg.get("PE") or 0
        if ref and (ce or pe):
            p["otype"] = "CE" if abs(ref - ce) <= abs(ref - pe) else "PE"
            p["otype_inferred"] = True
            p.pop("otype_unclear", None)
    return positions


def build_workbook(path: str, positions: list[dict], chains: dict[str, dict], cfg: dict) -> dict:
    """cfg keys: symbol, step, span, spot. One Sheet1, one expiry block per page."""
    open_positions = resolve_option_types([p for p in positions if p["qty"]], chains)
    expiries = sorted({p["expiry"] for p in open_positions if p.get("expiry")})

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.column_dimensions["A"].width = WIDTH_STRIKE
    ws.column_dimensions["B"].width = WIDTH_QTY
    ws.column_dimensions["C"].width = WIDTH_PREMIUM

    blocks, row = [], 2
    for expiry in expiries:
        book = [p for p in open_positions if p.get("expiry") == expiry]
        info = _block(ws, row, book, (chains or {}).get(expiry) or {}, cfg, expiry)
        if not info:
            continue
        for col in range(4, 4 + info["scenarios"]):
            ws.column_dimensions[get_column_letter(col)].width = WIDTH_SCENARIO
        blocks.append(dict(info, expiry=expiry))
        ws.row_breaks.append(Break(id=info["last"]))     # next expiry starts on a fresh page
        row = info["last"] + 3

    if blocks:
        widest = max(b["scenarios"] for b in blocks)
        ws.print_area = f"A1:{get_column_letter(3 + widest)}{blocks[-1]['last']}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0                       # 1 page wide; each expiry block is its own page
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.sheet_view.zoomScale = 100

    wb.save(path)
    return {
        "path": path,
        "positions": len(open_positions),
        "skipped_closed": len(positions) - len(open_positions),
        "blocks": len(blocks),
        "strike_rows": sum(b["rows"] for b in blocks),
        "scenarios": max((b["scenarios"] for b in blocks), default=0),
        "expiries": [b["expiry"] for b in blocks],
        "live_expiries": sorted(chains or {}),
        "spot": blocks[0]["spot"] if blocks else None,
    }
